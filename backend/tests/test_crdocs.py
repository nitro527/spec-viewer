import zipfile

from docx import Document

from specview import crdocs


def _cr_docx(path):
    d = Document()
    t = d.add_table(rows=2, cols=7)
    t.cell(0, 0).text = "CHANGE REQUEST"
    for i, v in enumerate(["38.211", "CR", "0123", "rev", "1", "Current version:", "16.0.0"]):
        t.cell(1, i).text = v
    d.add_table(rows=1, cols=2).cell(0, 0).text = "Proposed change affects:"
    t = d.add_table(rows=7, cols=4)
    rows = [("Title:", "Corrections to NR-U"), ("Source to WG:", "Ericsson"),
            ("Category:", "F"), ("Reason for change:", "Misalignment with RAN2."),
            ("Summary of change:", "Fix the LBT parameter."), ("Clauses affected:", "4.3.1, Table 7.4.1.5.3-1")]
    for r, (k, v) in enumerate(rows):
        t.cell(r, 0).text, t.cell(r, 1).text = k, v
    c = t.rows[6].cells
    c[0].text, c[1].text, c[2].text, c[3].text = "Work item code:", "NR_unlic-Core", "Date:", "2020-02-28"
    d.save(path)


def test_meeting_folder():
    assert crdocs.meeting_folder("RP-200186", "RAN#87-e").endswith("/tsg_ran/TSG_RAN/TSGR_87e")
    assert crdocs.meeting_folder("RP-242929", "RAN#106").endswith("/TSGR_106")
    assert crdocs.meeting_folder("R1-1708219", "RAN1#89") is None


def test_parse_clauses():
    assert crdocs.parse_clauses("5.3.1, Table 7.4.1.5.3-1, Annex B") == ["5.3.1", "7.4.1.5.3", "B"]


def test_pack_extract_and_cover(tmp_path):
    doc = tmp_path / "38211_CR0123_(Rel-16)_ R1-2001234.docx"
    _cr_docx(doc)
    pack = tmp_path / "RP-200186.zip"
    with zipfile.ZipFile(pack, "w") as z:
        z.write(doc, doc.name)
        z.writestr("38213_CR0099_(Rel-16)_ R1-2009999.docx", b"other spec")
        z.writestr("RP-200186_cover.doc", b"cover")
    got = crdocs.find_cr_doc(pack, "38.211", "0123", tmp_path / "out")
    cover = crdocs.parse_cover(got)
    assert (cover.spec, cover.cr, cover.rev, cover.current_version) == ("38.211", "0123", "1", "16.0.0")
    assert cover.title == "Corrections to NR-U" and cover.category == "F"
    assert cover.wi_code == "NR_unlic-Core" and cover.date == "2020-02-28"
    assert cover.reason == "Misalignment with RAN2."
    assert cover.clauses == ["4.3.1", "7.4.1.5.3"]
    assert cover.wg_tdoc == "R1-2001234"


def test_meeting_folder_old_and_new_names():
    assert crdocs.meeting_folder("RP-200186", "RAN#87-e").endswith("/TSGR_87e")
    assert crdocs.meeting_folder("RP-091168", "RAN_46").endswith("/TSGR_46")
    assert crdocs.meeting_folder("RP-1", "TSG RAN #82").endswith("/TSGR_82")


def test_old_pack_with_doc_cr(tmp_path, monkeypatch):
    from specview import download

    converted = []

    def fake_convert(src):
        converted.append(src.name)
        _cr_docx(src.with_suffix(".docx"))
        return src.with_suffix(".docx")

    monkeypatch.setattr(download, "_convert_doc", fake_convert)
    pack = tmp_path / "RP-091168.zip"
    with zipfile.ZipFile(pack, "w") as z:
        z.writestr("RP-091168.doc", b"cover")
        z.writestr("36211_CR0142_(Rel-8)_R1-094495.doc", b"old word")
        z.writestr("36212_CR0094_(Rel-8)_R1-094119.doc", b"other spec")
    got = crdocs.find_cr_doc(pack, "36.211", "142", tmp_path / "out")
    assert got.suffix == ".docx" and converted == ["36211_CR0142_(Rel-8)_R1-094495.doc"]
    assert crdocs.find_cr_doc(pack, "36.211", "142", tmp_path / "out") == got and len(converted) == 1  # cached


def test_ttcn_cr_tdoc_found_by_year(tmp_path):
    # 38.523-3 lists RAN5 TTCN e-mail Tdocs (R5s...), filed under TTCN_CRs/<year>/Docs
    asked = []

    class Resp:
        status_code, content = 200, b"PK\x03\x04"

    class Session:
        def get(self, url, timeout=None):
            asked.append(url)
            return Resp()

    assert crdocs.fetch_pack("R5s260219", "RAN#113", tmp_path, Session()).name == "R5s260219.zip"
    assert asked == [f"{crdocs.FTP}/tsg_ran/WG5_Test_ex-T1/TTCN/TTCN_CRs/2026/Docs/R5s260219.zip"]


def test_cr_in_nested_zip_found_by_cover_number(tmp_path):
    # RAN4 packs: one zip per CR (with JSON band combinations), names without the CR number or even the spec
    doc = tmp_path / "Big CR on DC_R20_HPUE.docx"
    _cr_docx(doc)
    inner = tmp_path / "R4-2606264.zip"
    with zipfile.ZipFile(inner, "w") as z:
        z.write(doc, doc.name)
        z.writestr("JSON Files/DC_n1A.json", "{}")
    pack = tmp_path / "RP-261384.zip"
    with zipfile.ZipFile(pack, "w") as z:
        z.write(inner, inner.name)
        z.writestr("R4-2607775.zip", b"not a zip")
        z.writestr("R4-2607776 CR 38.211 removing tables.docx", b"another CR, unreadable")
        z.writestr("RP-261384_cover.docx", b"cover")
    got = crdocs.find_cr_doc(pack, "38.211", "123", tmp_path / "out")
    assert got.name == doc.name and crdocs.parse_cover(got).cr == "0123"
    assert crdocs.find_cr_doc(pack, "38.211", "124", tmp_path / "out2") is None


def test_single_cr_tdoc_with_macos_junk_and_coverage_in_name(tmp_path):
    doc = tmp_path / "R5-265096_38521-2_Spherical_Coverage_Test_Procedure.docx"
    _cr_docx(doc)
    pack = tmp_path / "R5-265096.zip"
    with zipfile.ZipFile(pack, "w") as z:
        z.write(doc, doc.name)
        z.writestr(f"__MACOSX/._{doc.name}", b"resource fork")
    assert crdocs.find_cr_doc(pack, "38.521-2", "1287", tmp_path / "out").name == doc.name


class _Site:
    """3GPP FTP stand-in: url -> text (listings) or zip bytes; anything else is a 404."""

    def __init__(self, pages):
        self.pages, self.asked = pages, []

    def get(self, url, timeout=None):
        self.asked.append(url)
        body = self.pages.get(url)

        class R:
            status_code = 200 if body is not None else 404
            content = body if isinstance(body, bytes) else (body or "").encode()
            text = body if isinstance(body, str) else ""
        return R()


def test_tsg_meeting_found_by_tdoc_number_when_change_history_has_none(tmp_path):
    ran = f"{crdocs.FTP}/tsg_ran/TSG_RAN"
    site = _Site({
        f"{ran}/": "".join(f'<a href="{ran}/{n}/">' for n in ("TSGR_95e", "TSGR_96", "TSGR_97e", "TSGR_AHs")),
        f"{ran}/TSGR_95e/Docs/": '<a href="RP-220001.zip">RP-220001.zip</a> RP-220950.zip',
        f"{ran}/TSGR_96/Docs/": "RP-221000.zip RP-221800.zip",
        f"{ran}/TSGR_97e/Docs/": "RP-222000.zip RP-223500.zip",
        f"{ran}/TSGR_96/Docs/RP-221713.zip": b"PK\x03\x04",
    })
    assert crdocs.fetch_pack("RP-221713", "", tmp_path, site).name == "RP-221713.zip"  # 36.323: no meeting given


def test_e_meeting_folder_tried_when_change_history_drops_the_e(tmp_path):
    ran = f"{crdocs.FTP}/tsg_ran/TSG_RAN"
    site = _Site({f"{ran}/TSGR_88e/Docs/RP-201167.zip": b"PK\x03\x04"})
    assert crdocs.fetch_pack("RP-201167", "RP-88", tmp_path, site).name == "RP-201167.zip"  # 36.302
