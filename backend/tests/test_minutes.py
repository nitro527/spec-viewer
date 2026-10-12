from docx import Document

from specview import minutes
from specview.pack import word_changes


def _report(path):
    d = Document()
    d.add_heading("Maintenance for Reference signals and QCL", level=4)
    d.add_paragraph("Agreement")
    d.add_paragraph("Text proposal for TS38.211")
    d.add_paragraph("6.4.1.1.1.2\tSequence generation when transform precoding is enabled")
    d.add_paragraph("- if sequenceGroupHopping equals 'enabled' and sequenceHopping equals 'disabled', "
                    "group hopping but not sequence hopping shall be used")
    d.add_heading("Maintenance for physical downlink control channel", level=4)
    d.add_paragraph("Agreement:")
    d.add_paragraph("The draft CR in R1-1814118 is endorsed (38.213)")
    d.add_paragraph("Offline to prepare a draft CR, see also R1-1814055")
    d.add_paragraph("Decision: The draft CR in R1-1814055 is endorsed (38.211)")
    d.save(path)


def test_parse_and_match(tmp_path):
    _report(tmp_path / "r.docx")
    items = minutes.parse_report(tmp_path / "r.docx")
    assert [(i["agenda"][:11], i["kind"]) for i in items] == [
        ("Maintenance", "Agreement"), ("Maintenance", "Agreement"), ("Maintenance", "Decision")]
    rep = {"meeting": "RAN1#94b", "url": "u", "items": items}
    old, new = "- if group hopping but not sequence hopping shall be used", "- if group hopping is enabled and sequence hopping is disabled"
    got = minutes.match_agreements({"text": new, "old_text": old, "changes": word_changes(old, new), "clause": "6.4.1.1.1.2"},
                                   [rep], set(), "38.211")
    assert got[0]["via"] == "clause" and "Reference signals" in got[0]["agenda"]
    got = minutes.match_agreements({"text": "x", "changes": [], "clause": "7.3.2.2"}, [rep], {"R1-1814055"}, "38.211")
    assert got[0]["text"].startswith("The draft CR in R1-1814055 is endorsed (38.211)")  # the decision about this Tdoc first


def test_labels_and_tdoc_order():
    assert minutes.meeting_label("https://x/WG1_RL1/TSGR1_100b_e") == "RAN1#100b-e"
    assert minutes.tdoc_value("R1-101711") < minutes.tdoc_value("R1-1813585") < minutes.tdoc_value("R1-2005128")


def _xlsx(rows):
    import io
    import zipfile

    strings, cells = [], []
    for r, row in enumerate(rows, 1):
        cs = []
        for c, v in enumerate(row):
            strings.append(v)
            cs.append(f'<c r="{chr(65 + c)}{r}" t="s"><v>{len(strings) - 1}</v></c>')
        cells.append(f'<row r="{r}">{"".join(cs)}</row>')
    ns = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("xl/sharedStrings.xml", f"<sst {ns}>" + "".join(f"<si><t>{s}</t></si>" for s in strings) + "</sst>")
        z.writestr("xl/worksheets/sheet1.xml", f"<worksheet {ns}><sheetData>{''.join(cells)}</sheetData></worksheet>")
    return buf.getvalue()


def test_tdoc_list_and_summaries():
    head = ["TDoc", "Title", "Source", "Type", "Agenda item", "Agenda item description", "TDoc Status", "Revised to"]
    rows = minutes.parse_tdoc_list(_xlsx([
        head,
        ["R1-1811838", "Summary of DMRS Issues", "Qualcomm", "discussion", "7.1.3", "Maintenance for Reference signals and QCL", "revised", "R1-1811997"],
        ["R1-1811997", "Summary of DMRS Issues", "Qualcomm", "discussion", "7.1.3", "Maintenance for Reference signals and QCL", "noted", ""],
        ["R1-1810565", "On DMRS sequence hopping", "ZTE", "discussion", "7.1.3", "Maintenance for Reference signals and QCL", "noted", ""],
        ["R1-1812427", "Feature lead summary for NR-U", "Nokia", "discussion", "7.2.2", "DL Signals and Channels", "noted", ""],
    ]))
    assert rows[0]["title"] == "Summary of DMRS Issues" and rows[0]["revised_to"] == "R1-1811997"
    got = minutes.feature_lead_summaries({"tdocs": rows}, "Maintenance for Reference signals and QCL")
    assert [t["tdoc"] for t in got] == ["R1-1811838", "R1-1811997"]  # every revision: v1 may hold the settled issue

    paras = ["Issue 3: DMRS sequence hopping",
             "[ZTE] There is no explicit agreement of no sequence and group hopping at the same time.",
             "", "", "", "Unrelated text about PT-RS density."]
    text, score = minutes.excerpt(paras, {"sequence", "group", "hopping"}, set())
    assert "no explicit agreement" in text and "PT-RS" not in text and score > 0
    assert minutes.excerpt(["PT-RS density per RB"], {"sequence", "group", "hopping"}, set()) == ("", 0)


def test_document_viewer_only_serves_meeting_documents(tmp_path):
    import pytest

    for folder, tdoc in [("https://evil.example/TSGR1_95", None), ("https://www.3gpp.org/ftp/tsg_ran/WG1_RL1/TSGR1_95", "../x"),
                         ("https://www.3gpp.org/ftp/tsg_ran/WG1_RL1/TSGR1_95/../../x", None)]:
        with pytest.raises(ValueError):
            minutes.document_html(folder, tdoc, tmp_path)


def test_a_draft_report_is_rechecked_but_kept_when_nothing_better_comes(tmp_path, monkeypatch):
    import json
    import os
    import time

    from specview import minutes

    folder = "https://www.3gpp.org/ftp/tsg_ran/WG1_RL1/TSGR1_120"
    cache = tmp_path / "minutes" / f"{minutes.folder_name(folder)}.json"
    cache.parent.mkdir(parents=True)
    draft = {"v": 2, "meeting": "RAN1#120", "folder": folder, "url": f"{folder}/Report/Draft_v010.zip",
             "items": [{"x": 1}], "tdocs": []}
    cache.write_text(json.dumps(draft))
    old = time.time() - 8 * 86400
    os.utime(cache, (old, old))

    def offline(*a, **k):
        raise minutes.Unreachable("down")

    monkeypatch.setattr(minutes, "report_url", offline)
    assert minutes.fetch_report(folder, tmp_path)["items"] == [{"x": 1}]  # offline: still the draft
    monkeypatch.setattr(minutes, "report_url", lambda f, s=None: None)
    assert minutes.fetch_report(folder, tmp_path)["items"] == [{"x": 1}]  # report gone: draft kept, timer reset
    assert json.loads(cache.read_text())["items"] == [{"x": 1}] and cache.stat().st_mtime > old + 86400
